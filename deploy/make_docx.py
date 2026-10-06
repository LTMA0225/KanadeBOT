"""生成《音游窝 QQ 查询机器人 · 部署教程》Word 文档。"""

import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\YOGA-S~1\AppData\Local\Temp\opencode\部署教程.docx"

doc = Document()

# 全局字体（中文：微软雅黑）
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def fix_heading_font(paragraph):
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def h(text, level=1):
    para = doc.add_heading(text, level=level)
    fix_heading_font(para)
    return para


def p(*parts):
    """段落；part 可以是 str 或 (文本, 是否加粗)。"""
    para = doc.add_paragraph()
    for part in parts:
        if isinstance(part, tuple):
            run = para.add_run(part[0])
            run.bold = part[1]
        else:
            para.add_run(part)
    return para


def ul(items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            run = para.add_run(item[0])
            run.bold = item[1]
        else:
            para.add_run(item)


# ============ 封面标题 ============
title = doc.add_heading("音游窝 QQ 查询机器人", 0)
fix_heading_font(title)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("部署教程  ·  v1.0")
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
doc.add_paragraph()

# ============ 一、项目简介 ============
h("一、项目简介", 1)
p("本机器人用于在 QQ 群 / 私聊中查询「音游窝」店内实时在店人员，附带娱乐互动指令。")

p(("功能一览：", True))
ul(
    [
        ("指令 /j", True),
        "：查询在店人员，发送「图片 + 文本」；",
        "　　图片：黑白宵崎奏主题，铺满整幅画布；玩家显示昵称、入店时间与游玩分钟数（每人一个小框）；管理员 / 士大夫仅显示在店状态；",
        "　　文本：玩家一行（英文逗号分隔）；有管理 / 士大夫时另起一行。",
        ("指令 /baka", True),
        "：回复 bakabaka。",
        "玩家数量为 0 时（只有管理 / 士大夫或无人）：回复「店内无玩家，快来吧唧！」。",
    ]
)

p(("技术组成：", True))
ul(
    [
        ("AstrBot", True),
        "：机器人框架（本插件运行于其中，提供图片渲染与消息收发）；",
        ("NapCat", True),
        "：QQ 协议端（OneBot v11，反向 WebSocket 接入 AstrBot）；",
        ("本插件 astrbot_plugin_dcs_whosin", True),
        "：数据抓取（自动登录站点）+ 名单整理 + 图片渲染。",
    ]
)

p(("部署包内容：", True))
ul(
    [
        "astrbot_plugin_dcs_whosin/ ——插件本体（含一键部署脚本与看门狗脚本）；",
        "部署教程.docx —— 本文档；",
        "快速开始.txt —— 极简步骤速查；",
        "示例图.png —— 查询出图效果示例。",
    ]
)

# ============ 二、准备工作 ============
h("二、准备工作", 1)
ul(
    [
        ("一台 Windows 10 / 11 电脑", True),
        "：建议 7×24 保持开机（机器人需常驻在线）；",
        ("Python 3.10 或更高版本", True),
        "：官网 www.python.org/downloads 下载安装，安装时务必勾选 “Add python.exe to PATH”；",
        ("机器人专用 QQ 号", True),
        "：建议使用独立小号（QQ 电脑版 NT 需已安装）；",
        ("音游窝站点账号", True),
        "：一个可以查看 /whosin 页面的账号（用户名 + 密码），用于机器人自动登录抓取数据；",
        ("稳定的网络连接", True),
        "：图片渲染需要联网。",
    ]
)
p(
    ("⚠ 账号安全提醒：", True),
    "请勿使用主号作为机器人；避免机器人短时间内高频收发消息、频繁进群退群（可能触发平台风控导致封号）。",
)

# ============ 三、部署步骤 ============
h("三、部署步骤（共 8 步）", 1)

h("第 1 步：安装 AstrBot", 2)
p("任选一种方式：")
ul(
    [
        ("方式 A（命令行，推荐）：", True),
        "打开终端（Win+R 输入 cmd），执行：  python -m pip install astrbot",
        "国内网络可加速：  python -m pip install -i https://pypi.tuna.tsinghua.edu.cn/simple astrbot",
        ("方式 B（桌面客户端）：", True),
        "下载安装 github.com/AstrBotDevs/AstrBot-desktop/releases（图形界面，适合新手）。",
    ]
)

h("第 2 步：部署本插件", 2)
p("将部署包中的 astrbot_plugin_dcs_whosin 整个文件夹复制到 AstrBot 的插件目录下：")
ul(
    [
        "命令行部署：在运行 astrbot 的目录下，放入 data/plugins/ ；",
        "桌面客户端：放入 %USERPROFILE%\\.astrbot\\data\\plugins\\ ；",
        "也可使用插件目录中的 deploy\\windows_install.bat 一键完成安装与部署（双击运行即可）。",
    ]
)

h("第 3 步：启动 AstrBot 并启用插件", 2)
ul(
    [
        "终端输入 astrbot 启动（桌面客户端直接打开）；",
        "浏览器访问 http://localhost:6185 （初始账号 astrbot，初始密码见启动日志，登录后请修改）；",
        "进入「插件」页，确认 astrbot_plugin_dcs_whosin 已加载并处于启用状态（若显示禁用请点击启用）。",
    ]
)

h("第 4 步：安装 NapCat 并登录机器人 QQ", 2)
ul(
    [
        "下载 NapCat.Shell.zip（github.com/NapNeko/NapCatQQ/releases），解压到任意目录（推荐本插件目录下的 NapCat 文件夹）；",
        "运行目录中的 launcher-user.bat（如无法注入，可改用管理员运行 launcher.bat）；",
        "浏览器访问 http://localhost:6099 （NapCat 管理面板），用手机 QQ 扫码登录机器人账号；",
        "提示：若此前登录过且会话未过期，可在面板中使用「快速登录」。",
    ]
)

h("第 5 步：配置 NapCat 反向连接（关键）", 2)
ul(
    [
        "打开 NapCat 面板 →「网络配置」→ 新建 →「WebSocket 反向」；",
        "URL 填写： ws://127.0.0.1:6199/ws ；",
        "保存并启用。",
    ]
)

h("第 6 步：配置 AstrBot 消息平台", 2)
ul(
    [
        "打开 AstrBot 面板 →「消息平台」→「新增」；",
        "类型选择 aiocqhttp（OneBot v11），保持默认（反向 WS 端口 6199），保存启用；",
        "启用后，AstrBot 日志出现「反向连接成功」即表示连接建立。",
    ]
)

h("第 7 步：配置插件", 2)
p("AstrBot 面板 →「插件」→ 音游窝查询机器人 →「配置」，重点填写：")
ul(
    [
        ("login_username", True),
        "：音游窝站点账号；",
        ("login_password", True),
        "：站点密码（保存后自动登录，会话失效会自动重登，无需手动维护 Cookie）；",
        "其他配置保持默认即可（目标地址、超时、图片开关、缓存等）。",
    ]
)

h("第 8 步：测试", 2)
ul(
    [
        "把机器人拉进群（或加好友私聊），发送： /j  ——应返回「图片 + 名单」；",
        "发送 /baka ——应回复 bakabaka；",
        ("重要：", True),
        "请务必用你自己的 QQ 发送命令；不要用机器人自己的账号给它发消息（机器人不会响应自己发的消息）；",
        "若群内不响应：确认 AstrBot 的唤醒前缀为 “/”（AstrBot 配置中可修改）。",
    ]
)

# ============ 四、持续运行 ============
h("四、持续运行建议（可选）", 1)
ul(
    [
        ("电源设置：", True),
        "设置 → 系统 → 电源和电池：接通电源时永不睡眠；合上盖子时不采取任何操作；",
        ("开机自启：", True),
        "Win+R 输入 shell:startup 打开启动文件夹，放入 AstrBot（或桌面客户端）与 NapCat 启动器的快捷方式；",
        ("看门狗（随插件提供）：", True),
        "deploy\\dcs-bot-watchdog.ps1 每 5 分钟自检：AstrBot / NapCat 掉线自动拉起、QQ 未登录自动快速登录、连接假死自动重连；",
        "　　注册方式（管理员或普通用户终端均可）：schtasks /Create /TN \"DCStreamBotWatchdog\" /TR \"powershell -NoProfile -ExecutionPolicy Bypass -File \"完整路径\\dcs-bot-watchdog.ps1\"\" /SC MINUTE /MO 5 /F",
        "临时停用自愈： schtasks /Change /TN \"DCStreamBotWatchdog\" /DISABLE （恢复用 /ENABLE）。",
    ]
)

# ============ 五、常见问题 ============
h("五、常见问题（FAQ）", 1)

p(("Q1：群里发 /j 完全没有反应？按顺序检查：", True))
ul(
    [
        "① 命令要带唤醒前缀：「/j」（不是「j」）；",
        "② AstrBot「插件」页确认插件处于启用状态；",
        "③ AstrBot「消息平台」确认 aiocqhttp 适配器已启用，且日志显示「反向连接成功」；",
        "④ NapCat 面板确认 QQ 处于已登录状态；",
        "⑤ 不要用机器人自己的账号测试（机器人不响应自己发的消息）。",
    ]
)

p(("Q2：提示「目标页面要求登录」？", True))
p("在插件配置中填写 login_username / login_password（站点账号密码）即可，插件会全自动登录与续期。")

p(("Q3：图片发送失败 / 显示不出图片？", True))
ul(
    [
        "确认电脑可正常联网（图片由 AstrBot 渲染服务生成）；",
        "建议 AstrBot 与 NapCat 部署在同一台电脑上。",
    ]
)

p(("Q4：机器人账号掉线了？", True))
ul(
    [
        "打开 NapCat 面板（http://localhost:6099）：若显示二维码则用手机 QQ 扫码重新登录；",
        "再次登录后可点击「快速登录」以便下次免扫码；",
        "插件目录中 NapCat\\config\\webui.json 的 autoLoginAccount 可设置为机器人 QQ 号，实现 NapCat 重启后自动恢复登录。",
    ]
)

p(("Q5：换电脑 / 迁移部署？", True))
p("将整个部署包复制到新电脑，按「三、部署步骤」重新操作一遍即可（插件配置也可以整体复制 AstrBot 的 data/config/astrbot_plugin_dcs_whosin_config.json）。")

# ============ 六、目录结构 ============
h("六、插件目录结构说明", 1)
p("astrbot_plugin_dcs_whosin/")
ul(
    [
        "main.py —— 插件入口（指令注册）；",
        "dcs_api.py —— 页面抓取、登录、名单与时长解析；",
        "dcs_render.py —— 图片渲染与缓存；",
        "templates/whosin.html —— 图片模板（可自行修改样式）；",
        "assets/logo.svg —— 标志素材；",
        "skills/ —— 随插件提供的 AI 技能说明；",
        "_conf_schema.json —— 可视化配置定义；",
        "metadata.yaml —— 插件元数据；",
        "requirements.txt —— Python 依赖；",
        "deploy/ —— 部署脚本（windows_install.bat / linux_install.sh / docker-compose.yml / dcs-bot-watchdog.ps1）。",
    ]
)

doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = end.add_run("—— 教程结束 ——")
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

doc.save(OUT)
print("DOCX saved:", OUT)
